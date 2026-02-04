#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
tickets_pipeline.py

Split raw Sopra HR tickets and classify support team (DSN / Outils / Appli) with priority:
  1) Product-code mapping from "Code Produit.xlsx" via H2 lines
     - Example H2: "H2 PAPFRINT PAYFR:900" -> product "PAYFR:900"
     - Supports exact ("PAYFR - 900") and wildcard ("API4YCATFPE - 2XX") Excel codes
  2) Weighted text-only classifier (explicit closing > H2 presence > module > keywords)
  3) Legacy Excel token fallback (if still Unknown)

Closing-line semantics (Option A):
  - Lines like: "\t\tCP\t3\tPAPFRINT" (status, level, teamcode)
  - Status codes are decoded to human explanations (stored to CSV/front-matter).
  - Final statuses (e.g., CP, CA, CN, CH, CU, C2, MC, C4, CO, CS) boost confidence ONLY; they never change the team.

Additionally:
  - 'system' field has been removed from outputs.
  - The pipeline watches both raw file and Excel file in --watch mode; Excel updates auto-reload AND re-process.

Author: ZAIDI Mohamed Yassine + Copilot
"""

import argparse
import csv
import hashlib
import re
import time
from collections import Counter, defaultdict
from pathlib import Path
from typing import Dict, List, Tuple

# -----------------------------------------------------------------------------
# Utilities
# -----------------------------------------------------------------------------

FORBIDDEN_WIN_CHARS_RE = re.compile(r'[<>:"/\\|?\*]')

def sanitize_filename(text: str, max_len: int = 150) -> str:
    text = (text or "").strip()
    text = FORBIDDEN_WIN_CHARS_RE.sub('_', text)
    text = re.sub(r'\s+', '', text)   # remove spaces
    text = re.sub(r'_+', '_', text)   # collapse multiple underscores
    return text[:max_len] or "no_name"

def block_hash(text: str) -> str:
    return hashlib.md5(text.encode('utf-8', errors='ignore')).hexdigest()

def normalize_key(text: str) -> str:
    return re.sub(r'[^A-Z0-9]', '', str(text).upper())

# -----------------------------------------------------------------------------
# Ticket boundary & metadata
# -----------------------------------------------------------------------------

# Closing status detector lines look like: "\t\tCP\t3\tPAPFRINT"
#   group(1)=status_code, group(2)=level, group(3)=teamcode
END_STATUS_RE = re.compile(r'^\t\t([A-Z]{2})\t(\d+)\t([A-Z0-9_]+)')

# Header detection
START_RE = re.compile(r'^\d+\t')                 # first column integer + TAB
REF_RE   = re.compile(r'^(FR|AF|SP)\sW\d{5,}$')  # e.g. "FR W210205" (strict)
SYSTEM_WHITELIST = {'UNO','NZO','NZR','AS2','UN6','0C2'}  # only for header guard

def is_valid_header(line: str) -> bool:
    if not START_RE.match(line):
        return False
    cols = line.split('\t')
    if len(cols) < 6:
        return False

    # 0: ticket_no, 1: site_env, 2: offer, 3: system, 4: reference, 5: title
    system    = cols[3].strip()
    reference = cols[4].strip()
    title     = cols[5].strip()

    if not REF_RE.match(reference):
        return False
    if not title or title.strip('*').strip() == '':
        return False

    offer = cols[2].strip()
    if offer and len(offer) > 20:
        return False
    if system and (system not in SYSTEM_WHITELIST) and len(system) > 24:
        return False
    return True

def detect_segments(lines: List[str]) -> List[Tuple[int, int]]:
    starts = [i for i, l in enumerate(lines) if is_valid_header(l)]
    return [(s, starts[i+1] if i+1 < len(starts) else len(lines)) for i, s in enumerate(starts)]

def parse_header(line: str) -> Dict[str, str]:
    cols = line.split('\t')
    while len(cols) < 6:
        cols.append('')
    meta = {
        "ticket_no": cols[0].strip(),
        "site_env":  cols[1].strip(),
        "offer":     cols[2].strip(),
        # "system":  cols[3].strip(),   # removed from outputs per request
        "reference": cols[4].strip(),
        "title":     cols[5].strip(),
    }
    return meta

# -----------------------------------------------------------------------------
# Team classification (Weighted Text-only heuristic)
# -----------------------------------------------------------------------------

APPLI_CODES  = {'PAPFRINT', 'PAYFR', 'PAYFG', 'PAYFC', 'PADFR', 'PADFG', 'PADFE', 'PAPIN', 'COLTER', 'IPS', 'PAYFRINT', 'PAPFR'}
OUTILS_CODES = {'DESGN', 'CSTC', 'HRCT', 'HRASPACE', 'HRCTV2', 'QRYBT', 'PROCREP'}
DSN_CODES    = {'DSN', 'REGDSN', 'ESPDSN'}

# H2 presence (legacy signal for the text classifier; product mapping uses dedicated parser)
H2_RE      = re.compile(r'\bH2\s+([A-Z0-9_]+)\b')
MODULE_RE  = re.compile(r'\b([A-Z]{3,10})(?::[A-Z0-9]+)?\b')

KW_DSN   = {'DSN', 'N4DS', 'URSSAF', 'SIREN', 'BLOC', 'DECLARATION', 'DADS-U'}
KW_APPLI = {'PAIE', 'PAIEMENT', 'CET', 'ABSENCE', 'CONGE', 'MALADIE', 'RUBRIQUE', 'OPPOSITION', 'NRB', 'NRA', 'BULLETIN', 'MANDATEMENT'}
KW_OUTIL = {'DESIGN CENTER', 'COMPILATION', 'GENERATION', 'GENERATION PHYSIQUE', 'BNA', 'BNK', 'DBA', 'DBI', 'HRCT', 'HRQUERY', 'QUERY', 'SPACE', 'SERVEUR'}

WEIGHTS = {
    'explicit_closing': 5,
    'h2_route': 4,
    'module': 2,
    'keyword': 1,
    'explicit_closing_unknown': 0,
    'h2_route_unknown': 0,
}

def classify_team_text_only(block: str) -> Tuple[str, str, float, Dict]:
    signals = []
    codes_detected = set()

    # 1) explicit closing teamcode as a strong hint (3rd token on closing lines)
    for m in END_STATUS_RE.finditer(block):
        teamcode = m.group(3).upper()
        codes_detected.add(teamcode)
        if teamcode in DSN_CODES:
            signals.append(('explicit_closing', 'DSN', teamcode))
        elif teamcode in OUTILS_CODES:
            signals.append(('explicit_closing', 'Outils', teamcode))
        elif teamcode in APPLI_CODES:
            signals.append(('explicit_closing', 'Appli', teamcode))
        else:
            signals.append(('explicit_closing_unknown', 'Unknown', teamcode))

    # 2) H2 presence
    for m in H2_RE.finditer(block):
        code = m.group(1).upper()
        codes_detected.add(code)
        if code in DSN_CODES or code.startswith('REGDSN') or code.startswith('ESPDSN'):
            signals.append(('h2_route', 'DSN', code))
        elif code in OUTILS_CODES:
            signals.append(('h2_route', 'Outils', code))
        elif code in APPLI_CODES:
            signals.append(('h2_route', 'Appli', code))
        else:
            signals.append(('h2_route_unknown', 'Unknown', code))

    # 3) Module hints
    for m in MODULE_RE.finditer(block):
        code = m.group(1).upper()
        if code in {'H2', 'CP', 'CI', 'CN', 'CQ', 'CU'}:
            continue
        if code in DSN_CODES or code.startswith('REGDSN') or code.startswith('ESPDSN'):
            signals.append(('module', 'DSN', code))
        elif code in OUTILS_CODES:
            signals.append(('module', 'Outils', code))
        elif code in APPLI_CODES:
            signals.append(('module', 'Appli', code))

    # 4) Keywords
    U = block.upper()
    if any(k in U for k in KW_DSN):
        signals.append(('keyword', 'DSN', 'kw_dsn'))
    if any(k in U for k in KW_APPLI):
        signals.append(('keyword', 'Appli', 'kw_appli'))
    if any(k in U for k in KW_OUTIL):
        signals.append(('keyword', 'Outils', 'kw_outil'))

    score = Counter()
    for kind, team, _ in signals:
        score[team] += WEIGHTS.get(kind, 0)

    if not score:
        return ('Unknown', 'unknown', 0.0, {'codes_detected': sorted(codes_detected), 'score': dict(score)})

    best_team, best_score = score.most_common(1)[0]
    total = sum(score.values()) or 1
    confidence = round(best_score / total, 3)

    source = 'unknown'
    for preferred in ('explicit_closing', 'h2_route', 'module', 'keyword'):
        if any(k == preferred and t == best_team for (k, t, _) in signals):
            source = preferred
            break

    return (best_team, source, confidence, {'codes_detected': sorted(codes_detected), 'score': dict(score)})

# -----------------------------------------------------------------------------
# Excel product mapping (exact + wildcard) → DSN/Outils/Appli
# -----------------------------------------------------------------------------

def is_wildcard_code_str(s: str) -> bool:
    return bool(re.search(r'\b[0-9]XX\b', s.upper()))

def excel_code_to_pattern(excel_code_text: str) -> Tuple[str, str]:
    """
    Excel 'Code' -> (normalized_exact_or_left, digit_prefix_or_None)
      'PAYFR - 900'       -> ('PAYFR900', None)
      'API4YCATFPE - 2XX' -> ('API4YCATFPE', '2')
    """
    text = str(excel_code_text or "").upper().strip()
    parts = [p.strip() for p in text.split('-')]
    if len(parts) == 2:
        left, right = parts
        left_norm  = normalize_key(left)
        right_norm = normalize_key(right)
        m = re.fullmatch(r'([0-9])XX', right.strip().upper())
        if m:
            return (left_norm, m.group(1))
        return (f"{left_norm}{right_norm}", None)
    return (normalize_key(text), None)

def domain_to_team(domain_upper: str) -> str:
    """
    Map Excel 'Domaine' (UPPER) to {DSN, Outils, Appli}
    """
    if not domain_upper:
        return "Unknown"
    if "DSN" in domain_upper:
        return "DSN"
    if any(k in domain_upper for k in ("CORE", "WEBTOOLS", "HRCT", "API", "RUNTOOLS", "STUDIO", "DELIVERY", "IPS")):
        return "Outils"
    return "Appli"

def load_excel_mapping(xlsx_path: Path) -> Dict[str, object]:
    """
    Returns:
      "_EXACT": normalized_code -> domain (UPPER)
      "_WILD" : list[(left_norm, digit_prefix, domain)]
    """
    try:
        from openpyxl import load_workbook
    except Exception as e:
        print(f"[WARN] Excel mapping disabled (openpyxl not available): {e}")
        return {"_EXACT": {}, "_WILD": []}

    if not xlsx_path or not xlsx_path.exists():
        print(f"[WARN] Excel file not found: {xlsx_path}")
        return {"_EXACT": {}, "_WILD": []}

    wb = load_workbook(filename=str(xlsx_path), read_only=True, data_only=True)
    ws = wb.active

    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    headers = {str(v).strip().lower(): i for i, v in enumerate(header_row) if v is not None}
    code_idx = headers.get("code")
    dom_idx  = headers.get("domaine") or headers.get("domain")
    if code_idx is None or dom_idx is None:
        print("[WARN] Could not find 'Code' and 'Domaine' columns in Excel.")
        return {"_EXACT": {}, "_WILD": []}

    mapping_exact: Dict[str, str] = {}
    mapping_wild : List[Tuple[str, str, str]] = []

    for row in ws.iter_rows(min_row=2, values_only=True):
        code_cell = str(row[code_idx] or "").strip()
        dom_cell  = str(row[dom_idx]  or "").strip().upper()
        if not code_cell:
            continue
        norm_or_left, digit_prefix = excel_code_to_pattern(code_cell)
        if digit_prefix is None:
            mapping_exact[norm_or_left] = dom_cell
        else:
            mapping_wild.append((norm_or_left, digit_prefix, dom_cell))

    return {"_EXACT": mapping_exact, "_WILD": mapping_wild}

# -----------------------------------------------------------------------------
# Product-code-based classifier (from H2 lines)
# -----------------------------------------------------------------------------

# Example H2 line: "... H2 PAPFRINT PAYFR:900"
H2_PRODUCT_RE = re.compile(r'\b(H\d)\s+([A-Z0-9_]+)\s+([A-Z0-9_]+):(\d{2,4})\b')

def normalize_product_code(code_word: str, code_num: str) -> str:
    return normalize_key(f"{code_word}{code_num}")  # "PAYFR", "900" -> "PAYFR900"

def classify_team_via_product_code(block: str, excel_maps: Dict[str, object]) -> Tuple[str, str, Dict]:
    """
    Use product code parsed from H2 lines (e.g., PAYFR:900) to lookup Excel 'Code' and infer DSN/Outils/Appli.
    Returns (team, domain, evidence)
    """
    if not excel_maps:
        return ("Unknown", "", {"reason": "no_excel"})

    m_exact: Dict[str, str] = excel_maps.get("_EXACT", {})
    m_wild : List[Tuple[str, str, str]] = excel_maps.get("_WILD", [])

    products = []
    for m in H2_PRODUCT_RE.finditer(block):
        word = m.group(3).upper()
        digits = m.group(4)
        norm = normalize_product_code(word, digits)  # e.g., "PAYFR900"
        products.append((word, digits, norm))

    if not products:
        return ("Unknown", "", {"reason": "no_product_code_in_h2"})

    # 1) exact match
    for (w, d, norm) in products:
        dom = m_exact.get(norm)
        if dom:
            return (domain_to_team(dom), dom, {"match": "exact", "product": f"{w}:{d}", "excel_key": norm})

    # 2) wildcard by left + first digit
    for (w, d, norm) in products:
        left_norm = normalize_key(w)
        first_digit = d[0]
        for (ex_left, ex_digit, dom) in m_wild:
            if ex_left == left_norm and ex_digit == first_digit:
                return (domain_to_team(dom), dom, {"match": "wildcard", "product": f"{w}:{d}", "excel_left": ex_left, "digit_prefix": ex_digit})

    return ("Unknown", "", {"reason": "no_excel_match", "products": [f"{w}:{d}" for w, d, _ in products]})

# -----------------------------------------------------------------------------
# Legacy Excel token fallback (used if Unknown after text classifier)
# -----------------------------------------------------------------------------

def classify_team_excel_fallback(block: str, mapping_exact: Dict[str, str]) -> Tuple[str, str, Dict]:
    if not mapping_exact:
        return ("Unknown", "", {"reason": "no_excel_mapping"})

    U = block.upper()
    tokens = set()
    for tok in ("HRCT","HRCTV2","HRASPACE","OPENHR","STUDIO","DSGXX","ADDIN","DMS","API","HRWEB",
                "DSN","DADSU","REGDSN","ESPDSN"):
        if tok in U:
            tokens.add(tok)

    if not tokens:
        if any(k in U for k in ("DSN","DADSU","DADS-U","REGDSN","ESPDSN")):
            return ("DSN", "DSN", {"reason": "broad_dsn_keyword"})
        return ("Unknown", "", {"reason": "no_tokens"})

    hits = Counter()
    evidence = defaultdict(list)
    for code_key_norm, dom in mapping_exact.items():
        for t in tokens:
            if normalize_key(t) in code_key_norm:
                hits[dom] += 1
                evidence[dom].append((t, code_key_norm))

    if not hits:
        if "DSN" in tokens or "DADSU" in tokens or "REGDSN" in tokens or "ESPDSN" in tokens:
            return ("DSN", "DSN", {"reason": "token_contains_dsn"})
        return ("Unknown", "", {"reason": "tokens_no_match"})

    best_domain, _ = hits.most_common(1)[0]
    return (domain_to_team(best_domain), best_domain, {"hits": dict(hits), "evidence": {k: list(v) for k, v in evidence.items()}})

# -----------------------------------------------------------------------------
# Closing status semantics (Option A) → explanation + confidence adjustment
# -----------------------------------------------------------------------------

# Explanations (from your screenshot list). Not exhaustive, but covers key ones.
CLOSING_STATUS_EXPLANATIONS = {
    "AA": "To be processed",
    "AF": "Preparation Hotfix delivery",
    "AI": "Information Received",
    "AK": "In Progress",
    "AW": "Waiting information",
    "CA": "New patch 4YOU (closed)",
    "CH": "Sent to Services (closed)",
    "CI": "Waiting for customer information",
    "CN": "New patch delivered (closed)",
    "CO": "Old patch delivered (closed)",
    "CP": "Request processed (closed)",
    "CQ": "Corrected in future release (micro modules)",
    "CR": "Corrected in existing release (micro modules)",
    "CS": "Bypass solution given (closed)",
    "CT": "Client part order",
    "CU": "Incorrect use (closed)",
    "CV": "Enhancement approved for future version",
    "CW": "Enhancement to study for future version",
    "CX": "Explicit abort by client (do not use)",
    "CY": "Implicit abort by client (do not use)",
    "C2": "Cancelled or service event (closed)",
    "MC": "In V or V1 only - Closed without solution"
}

# Final (closed) statuses that boost confidence (Option A).
# Included CO and CS because their labels indicate (closed).
STATUS_FINAL = {"CA","CN","CP","CH","CU","C2","MC","C4","CO","CS"}

def explain_status(code: str) -> str:
    return CLOSING_STATUS_EXPLANATIONS.get(code.upper(), "")

def adjust_confidence_by_status(conf: float, observed_statuses: List[str]) -> float:
    """
    Option A: If any final (closed) status appears, raise confidence to at least 0.88.
    Never changes the team.
    """
    if any(s in STATUS_FINAL for s in observed_statuses):
        return max(conf, 0.88)
    return conf

# -----------------------------------------------------------------------------
# Index I/O (incremental UPSERT by 'reference')
# -----------------------------------------------------------------------------

INDEX_FIELDS = [
    'ordinal','ticket_no','site_env','offer','reference','title',
    'start_line','end_line',
    'closing_status_line','closing_status_code','closing_status_explanation','closing_level','closing_teamcode',
    'file_name','support_team','team_source','confidence','content_hash'
]

def load_existing_index(index_csv: Path) -> Dict[str, Dict[str, str]]:
    if not index_csv.exists():
        return {}
    out = {}
    with index_csv.open('r', encoding='utf-8') as f:
        reader = csv.DictReader(f)
        for row in reader:
            ref = (row.get('reference') or '').strip()
            if ref:
                out[ref] = row
    return out

def write_index_incremental(index_csv: Path, new_rows: List[Dict[str, str]]):
    """
    Incrementally upsert rows into the CSV by 'reference'.
    - Keeps all existing rows.
    - Replaces rows that share the same 'reference' as any of new_rows.
    - Appends new references.
    - Writes atomically via a temporary file.
    """
    existing_by_ref: Dict[str, Dict[str, str]] = {}
    if index_csv.exists():
        with index_csv.open('r', encoding='utf-8') as f:
            reader = csv.DictReader(f)
            for row in reader:
                ref = (row.get('reference') or '').strip()
                if ref:
                    existing_by_ref[ref] = row

    for r in new_rows:
        ref = (r.get('reference') or '').strip()
        if not ref:
            ref = f"__NOREF__:{r.get('ordinal','')}"
        existing_by_ref[ref] = r

    temp_path = index_csv.with_suffix(index_csv.suffix + ".tmp")
    index_csv.parent.mkdir(parents=True, exist_ok=True)
    with temp_path.open('w', newline='', encoding='utf-8') as f:
        writer = csv.DictWriter(f, fieldnames=INDEX_FIELDS)
        writer.writeheader()

        def sort_key(item):
            row = item[1]
            try:
                return (int(row.get('ordinal', '0')), row.get('reference',''))
            except ValueError:
                return (0, row.get('reference',''))

        for _, row in sorted(existing_by_ref.items(), key=sort_key):
            writer.writerow(row)

    temp_path.replace(index_csv)

# -----------------------------------------------------------------------------
# Core processing
# -----------------------------------------------------------------------------

def process_raw_file(raw_file: Path, out_dir: Path, index_csv: Path, excel_maps: Dict[str, object]) -> Tuple[int, int, int]:
    out_dir.mkdir(parents=True, exist_ok=True)

    text = raw_file.read_text(encoding='utf-8', errors='ignore')
    lines = text.splitlines()
    segments = detect_segments(lines)
    existing = load_existing_index(index_csv)

    rows: List[Dict[str, str]] = []
    created = updated = 0

    for ordinal, (s, e) in enumerate(segments, start=1):
        header = lines[s]
        meta = parse_header(header)
        reference = meta['reference']
        title     = meta['title']
        block_text = '\n'.join(lines[s:e])
        content_md5 = block_hash(block_text)

        # Closing status capture
        closing_line = ''
        closing_code = ''
        closing_level = ''
        closing_teamcode = ''
        closing_expl = ''
        status_codes_seen: List[str] = []

        for k in range(s, e):
            m = END_STATUS_RE.match(lines[k])
            if m:
                closing_line = lines[k].strip()
                closing_code = m.group(1).upper()
                closing_level = m.group(2)
                closing_teamcode = m.group(3).upper()
                status_codes_seen.append(closing_code)
                closing_expl = explain_status(closing_code)

        # --- PRIORITY ORDER ---
        # 0) Product code via Excel (H2 lines) → excel_product
        team = src = ""
        conf = 0.0
        evidence: Dict = {}

        team_prod, dom_prod, ev_prod = classify_team_via_product_code(block_text, excel_maps)
        if team_prod != "Unknown":
            team, src, conf = team_prod, "excel_product", 0.90
            evidence = {"excel_domain": dom_prod, "excel_product_evidence": ev_prod}
        else:
            # 1) Text-only classifier
            team_txt, src_txt, conf_txt, evid_txt = classify_team_text_only(block_text)
            team, src, conf, evidence = team_txt, src_txt, conf_txt, evid_txt

            # 2) Legacy Excel token fallback (only if Unknown)
            if team == "Unknown":
                team_xl, domain_xl, ev_xl = classify_team_excel_fallback(block_text, excel_maps.get("_EXACT", {}))
                if team_xl != "Unknown":
                    team, src, conf = team_xl, "excel", 0.75
                    evidence = {"excel_domain": domain_xl, "excel_evidence": ev_xl}

        # Option A: closing status → boost confidence only
        conf = adjust_confidence_by_status(conf, status_codes_seen)

        # Per-ticket file name
        base_ref = reference or f'noRef_{ordinal:04d}'
        ref_clean = sanitize_filename(base_ref)
        filename = f"{ordinal:04d}_{ref_clean}.txt"
        file_path = out_dir / filename

        # Incremental write
        prev = existing.get(reference)
        need_write = False
        if prev is None:
            need_write = True
            created += 1
        else:
            if prev.get('content_hash') != content_md5 or prev.get('file_name') != filename:
                need_write = True
                updated += 1

        # Write per-ticket file with extended front-matter
        if need_write:
            header_meta = (
                f"---\n"
                f"reference: {reference}\n"
                f"title: {title}\n"
                f"support_team: {team}\n"
                f"team_source: {src}\n"
                f"confidence: {conf}\n"
                f"content_hash: {content_md5}\n"
                f"closing_status_code: {closing_code}\n"
                f"closing_status_explanation: {closing_expl}\n"
                f"closing_level: {closing_level}\n"
                f"closing_teamcode: {closing_teamcode}\n"
                f"---\n"
            )
            file_path.write_text(header_meta + block_text.rstrip() + "\n", encoding='utf-8')

        # Row for CSV (note: 'system' removed)
        rows.append({
            'ordinal': str(ordinal),
            'ticket_no': meta['ticket_no'],
            'site_env':  meta['site_env'],
            'offer':     meta['offer'],
            'reference': reference,
            'title':     title,
            'start_line': str(s + 1),
            'end_line':   str(e),
            'closing_status_line': closing_line,
            'closing_status_code': closing_code,
            'closing_status_explanation': closing_expl,
            'closing_level': closing_level,
            'closing_teamcode': closing_teamcode,
            'file_name':  filename,
            'support_team': team,
            'team_source':  src,
            'confidence':   f"{conf:.3f}",
            'content_hash': content_md5
        })

    # 🔁 UPSERT (keeps previous rows; updates/append current ones)
    write_index_incremental(index_csv, rows)
    return (len(segments), created, updated)

# -----------------------------------------------------------------------------
# Watch loop with Excel auto-reload (and actual re-process after change)
# -----------------------------------------------------------------------------

def main():
    ap = argparse.ArgumentParser(description="Split raw tickets and classify DSN/Outils/Appli using product codes (Excel), with text and Excel fallbacks. Closing statuses add explanations and boost confidence (Option A).")
    ap.add_argument("--raw-file",      type=Path, required=True, help="Path to raw tickets export .txt")
    ap.add_argument("--out-dir",       type=Path, default=Path("data/output"), help="Output folder for per-ticket files")
    ap.add_argument("--index-csv",     type=Path, default=Path("data/tickets_index.csv"), help="CSV index (will be upserted)")
    ap.add_argument("--product-excel", type=Path, default=Path("data/Code Produit.xlsx"), help="Excel file with 'Code' and 'Domaine' columns")
    ap.add_argument("--watch", action="store_true", help="Watch raw file and Excel for changes and re-process automatically")
    ap.add_argument("--interval", type=int, default=30, help="Watch mode: polling interval (seconds)")
    args = ap.parse_args()

    # initial load
    excel_maps = load_excel_mapping(args.product_excel) if args.product_excel else {"_EXACT": {}, "_WILD": []}

    last_raw_mtime = None
    last_xls_mtime = args.product_excel.stat().st_mtime if args.product_excel and args.product_excel.exists() else None

    while True:
        try:
            raw_mtime = args.raw_file.stat().st_mtime
            xls_mtime = args.product_excel.stat().st_mtime if args.product_excel and args.product_excel.exists() else None

            # Detect Excel change BEFORE updating last_xls_mtime
            excel_changed = False
            if args.product_excel and xls_mtime is not None:
                if last_xls_mtime is None or xls_mtime != last_xls_mtime:
                    excel_changed = True
                    excel_maps = load_excel_mapping(args.product_excel)  # re-parse updated codes
                    print(f"[watch] Detected change in Excel: {args.product_excel}. Reloaded mapping.")

            raw_changed = (last_raw_mtime is None) or (raw_mtime != last_raw_mtime)

            # Re-process if raw file changed OR Excel changed
            if raw_changed or (args.watch and excel_changed):
                total, created, updated = process_raw_file(args.raw_file, args.out_dir, args.index_csv, excel_maps)
                print(f"[pipeline] tickets={total}  created={created}  updated={updated}  out={args.out_dir}  index={args.index_csv}")
                last_raw_mtime = raw_mtime
                last_xls_mtime = xls_mtime

            if not args.watch:
                break

            time.sleep(args.interval)
        except KeyboardInterrupt:
            print("Stopped.")
            break

if __name__ == "__main__":
    main()